#!/usr/bin/env python3
"""Build the NHL player-prop tracker workbook (goals, points, shots, parlays) from daily.py outputs.
Usage: python3 build_xlsx.py OUT_DIR WORKBOOK_PATH
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT, PATH = sys.argv[1], sys.argv[2]
J = lambda n: json.load(open(f'{OUT}/{n}')) if os.path.exists(f'{OUT}/{n}') else []
today, tpts, tsog, upcoming = J('today.json'), J('today_pts.json'), J('today_sog.json'), J('upcoming.json')
tracker, changes, acc, cal, top = J('tracker_view.json'), J('changelog_view.json'), J('acc.json'), J('cal.json'), J('top.json')
ptoday, plog, pbt, summary = J('parlays_today.json'), J('parlay_log.json'), J('parlay_bt.json'), J('summary.json')
champs = summary.get('champions', {})

F = 'Arial'
base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True); title = Font(name=F, size=14, bold=True)
muted = Font(name=F, size=9, italic=True, color='666666'); head = Font(name=F, size=10, bold=True, color='FFFFFF')
inp = Font(name=F, size=10, color='0000FF'); hfill = PatternFill('solid', fgColor='1F3A5F'); yfill = PatternFill('solid', fgColor='FFFF00')
howto = Font(name=F, size=10, bold=True, color='9C5700')
thin = Border(bottom=Side(style='thin', color='D0D0D0'))
green = PatternFill('solid', fgColor='D9EAD3'); red = PatternFill('solid', fgColor='F4CCCC'); grey = PatternFill('solid', fgColor='EFEFEF')
FAIR = '=IF({c}="","",IF({c}<0.5,"+"&ROUND(100*(1-{c})/{c},0),"-"&ROUND(100*{c}/(1-{c}),0)))'
IMPL = '=IF({o}="","",IF({o}>0,100/({o}+100),-{o}/(-{o}+100)))'
PCT, PCTD = '0.0%', '+0.0%;-0.0%;0.0%'
d0 = (today[0]['date'] if today else summary.get('run_date', ''))

wb = Workbook()

def table(ws, r0, headers, rows, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h); c.font = head; c.fill = hfill; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.font = base; c.border = thin
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + len(rows)

def header_block(ws, t, sub):
    ws['A1'] = t; ws['A1'].font = title
    ws['A2'] = sub; ws['A2'].font = muted

def odds_cols(ws, r0, last, pcell, ocol, icol, ecol, bcol):
    """Your-odds input + implied + edge + bet flag, for rows r0+1..last."""
    for i in range(r0 + 1, last + 1):
        ws[f'{ocol}{i}'].fill = yfill; ws[f'{ocol}{i}'].font = inp; ws[f'{ocol}{i}'].number_format = '+0;-0'
        ws[f'{icol}{i}'] = IMPL.format(o=f'{ocol}{i}'); ws[f'{icol}{i}'].number_format = PCT
        ws[f'{ecol}{i}'] = f'=IF({ocol}{i}="","",{pcell(i)}-{icol}{i})'; ws[f'{ecol}{i}'].number_format = PCTD
        ws[f'{bcol}{i}'] = f'=IF({ocol}{i}="","",IF({ecol}{i}>0,"Value","Pass"))'
    if last > r0:
        ws.conditional_formatting.add(f'{bcol}{r0+1}:{bcol}{last}', CellIsRule(operator='equal', formula=['"Value"'], fill=green))
        ws.conditional_formatting.add(f'{bcol}{r0+1}:{bcol}{last}', CellIsRule(operator='equal', formula=['"Pass"'], fill=red))

HOW = 'How to use: type your sportsbook\'s American odds in the yellow "Your odds" column (e.g. +135 → 135, -110 → -110). Edge and Bet? fill in automatically.'

# ---------------- Goals today
ws = wb.active; ws.title = 'Goals today'
header_block(ws, f'Anytime goalscorers — {d0}', f"Goal model {champs.get('goal', '')} · data through {summary.get('data_through', '')} · * = goalie not confirmed")
ws['A3'] = HOW; ws['A3'].font = howto
ws['A4'] = 'Example: P(goal) 46.0% with your odds 135 → book implied 42.6% → edge +3.4% → Value'; ws['A4'].font = muted
R0 = 6
last = table(ws, R0, ['Rank', 'Player', 'Team', 'Opp', 'P(goal)', 'Fair odds', 'TOI', 'PP TOI', 'xG/60', 'SOG/gm', 'Finish', 'DUE', 'Opp goalie',
                      'Your odds', 'Book implied', 'Edge', 'Bet?'],
             [[i + 1, r['player'], r['team'], r['opp'], r['p_goal'], None, r['toi'], r['pp_toi'], r['xg60'], r['sog'], r['finishing'], r['due'],
               r['opp_goalie'], None, None, None, None] for i, r in enumerate(today)],
             [6, 22, 6, 8, 9, 9, 6, 7, 7, 7, 7, 7, 22, 10, 11, 9, 8])
for i in range(R0 + 1, last + 1):
    ws[f'E{i}'].number_format = PCT; ws[f'F{i}'] = FAIR.format(c=f'E{i}')
    for col in 'GHJ': ws[f'{col}{i}'].number_format = '0.0'
    for col in 'IKL': ws[f'{col}{i}'].number_format = '0.00'
odds_cols(ws, R0, last, lambda i: f'E{i}', 'N', 'O', 'P', 'Q')
ws.freeze_panes = f'C{R0+1}'
ws['L6'].comment = Comment('Finishing-adjusted xG minus goals this season. Shown for reference; it did not predict goals in testing.', 'Model')

# ---------------- Points today
ws = wb.create_sheet('Points today')
header_block(ws, f'1+ point — {d0}', f"Points model {champs.get('point', '')} · chance the player records at least one goal or assist")
ws['A3'] = HOW; ws['A3'].font = howto
last = table(ws, R0, ['Rank', 'Player', 'Team', 'Opp', 'P(1+ point)', 'Fair odds', 'Points/gm', 'TOI', 'PP TOI', 'Opp goalie',
                      'Your odds', 'Book implied', 'Edge', 'Bet?'],
             [[i + 1, r['player'], r['team'], r['opp'], r['p_point'], None, r['pts_pg'], r['toi'], r['pp_toi'], r['opp_goalie'], None, None, None, None]
              for i, r in enumerate(tpts)], [6, 22, 6, 8, 11, 9, 10, 6, 7, 22, 10, 11, 9, 8])
for i in range(R0 + 1, last + 1):
    ws[f'E{i}'].number_format = PCT; ws[f'F{i}'] = FAIR.format(c=f'E{i}')
    ws[f'G{i}'].number_format = '0.00'; ws[f'H{i}'].number_format = '0.0'; ws[f'I{i}'].number_format = '0.0'
odds_cols(ws, R0, last, lambda i: f'E{i}', 'K', 'L', 'M', 'N')
ws.freeze_panes = f'C{R0+1}'

# ---------------- Shots today
ws = wb.create_sheet('Shots today')
header_block(ws, f'Shots on goal — {d0}', f"Shots model {champs.get('sog', '')} · projected shots and the chance of reaching each line")
ws['A3'] = 'How to use: type the line you are betting (2, 3 or 4 = "over 1.5 / 2.5 / 3.5") in "Line" and your odds in "Your odds". Edge and Bet? fill in.'
ws['A3'].font = howto
last = table(ws, R0, ['Rank', 'Player', 'Team', 'Opp', 'Proj SOG', 'P(2+)', 'Fair 2+', 'P(3+)', 'Fair 3+', 'P(4+)', 'Fair 4+', 'SOG/gm', 'TOI',
                      'Line', 'Your odds', 'Book implied', 'Edge', 'Bet?'],
             [[i + 1, r['player'], r['team'], r['opp'], r['sog_proj'], r['p_sog2'], None, r['p_sog3'], None, r['p_sog4'], None, r['sog'], r['toi'],
               None, None, None, None, None] for i, r in enumerate(tsog)], [6, 22, 6, 8, 9, 8, 8, 8, 8, 8, 8, 8, 6, 6, 10, 11, 9, 8])
for i in range(R0 + 1, last + 1):
    ws[f'E{i}'].number_format = '0.00'
    for p, fcol in (('F', 'G'), ('H', 'I'), ('J', 'K')):
        ws[f'{p}{i}'].number_format = PCT; ws[f'{fcol}{i}'] = FAIR.format(c=f'{p}{i}')
    ws[f'L{i}'].number_format = '0.0'; ws[f'M{i}'].number_format = '0.0'
    ws[f'N{i}'].fill = yfill; ws[f'N{i}'].font = inp
odds_cols(ws, R0, last, lambda i: f'IF(AND(N{i}>=2,N{i}<=4),CHOOSE(N{i}-1,F{i},H{i},J{i}),NA())', 'O', 'P', 'Q', 'R')
for i in range(R0 + 1, last + 1):   # blank edge until a valid line is typed
    ws[f'Q{i}'] = f'=IF(OR(O{i}="",N{i}="",N{i}<2,N{i}>4),"",CHOOSE(N{i}-1,F{i},H{i},J{i})-P{i})'
    ws[f'R{i}'] = f'=IF(Q{i}="","",IF(Q{i}>0,"Value","Pass"))'
ws.freeze_panes = f'C{R0+1}'

# ---------------- Parlays today
ws = wb.create_sheet('Parlays today')
header_block(ws, f'Model parlays — {d0}', 'Built from the top of each board, one player per game (teammate pair: two players on the same team). Paper-tracked every night.')
ws['A3'] = 'Type your sportsbook\'s parlay odds in the yellow column (e.g. +450 → 450). Each leg\'s house edge stacks, so expect most book prices to show Pass.'
ws['A3'].font = howto
last = table(ws, R0, ['Type', 'Legs', 'Leg chances', 'P(all hit)', 'Fair odds', 'Your odds', 'Book implied', 'Edge', 'Bet?'],
             [[p['type'], p['parlay'], p['leg_probs'], p['p_parlay'], None, None, None, None, None] for p in ptoday],
             [18, 80, 22, 11, 11, 11, 12, 9, 8])
for i in range(R0 + 1, last + 1):
    ws[f'D{i}'].number_format = PCT; ws[f'E{i}'] = FAIR.format(c=f'D{i}')
odds_cols(ws, R0, last, lambda i: f'D{i}', 'F', 'G', 'H', 'I')
if not ptoday: ws[f'A{R0+1}'] = 'No parlays today (not enough games).'; ws[f'A{R0+1}'].font = muted

# ---------------- Upcoming
ws = wb.create_sheet('Upcoming')
header_block(ws, 'Next days — early look (goals)', 'Goalies are guesses and lineups can change; each morning\'s run replaces these with a fresh board.')
last = table(ws, 4, ['Date', 'Player', 'Team', 'Opp', 'P(goal)', 'Fair odds', 'Opp goalie'],
             [[u['date'], u['player'], u['team'], u['opp'], u['p_goal'], None, u['opp_goalie']] for u in upcoming], [12, 22, 6, 8, 9, 9, 22])
for i in range(5, last + 1):
    ws[f'E{i}'].number_format = PCT; ws[f'F{i}'] = FAIR.format(c=f'E{i}')
ws.freeze_panes = 'C5'

# ---------------- Tracker
ws = wb.create_sheet('Tracker')
header_block(ws, 'Pick tracker', 'The model\'s top ten each night in each market, graded the next morning from NHL boxscores. "Did not play" picks are voided.')
TR0 = 4
last = table(ws, TR0, ['Date', 'Market', 'Rank', 'Player', 'Team', 'Opp', 'P(hit)', 'Fair odds', 'Model', 'Result', 'Hit'],
             [[t['date'], t['market'], t['rank'], t['player'], t['team'], t['opp'], t['p'], None, t['model'], t['result'], None] for t in tracker],
             [12, 10, 6, 22, 6, 8, 9, 9, 8, 13, 6])
for i in range(TR0 + 1, last + 1):
    ws[f'G{i}'].number_format = PCT; ws[f'H{i}'] = FAIR.format(c=f'G{i}')
    ws[f'K{i}'] = f'=IF(J{i}="Hit",1,IF(J{i}="Miss",0,""))'
TRL = max(last, TR0 + 1)
if last > TR0:
    rng = f'J{TR0+1}:J{last}'
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Hit"'], fill=green))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Miss"'], fill=red))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Pending"'], fill=grey))
ws.freeze_panes = f'E{TR0+1}'

# ---------------- Daily results (formulas over Tracker)
ws = wb.create_sheet('Daily results')
header_block(ws, 'Results', 'Expected hits = sum of the model\'s probabilities. Over many nights, Hits should track Expected if the model is honest.')
A_, M_, R_, P_ = (f'Tracker!${c}${TR0+1}:${c}${TRL}' for c in 'ABJG')
def stats_row(ws, i, crit):
    ws[f'C{i}'] = f'=COUNTIFS({crit},{R_},"Hit")+COUNTIFS({crit},{R_},"Miss")'
    ws[f'D{i}'] = f'=COUNTIFS({crit},{R_},"Hit")'
    ws[f'E{i}'] = f'=SUMIFS({P_},{crit},{R_},"Hit")+SUMIFS({P_},{crit},{R_},"Miss")'
    ws[f'F{i}'] = f'=IF(C{i}=0,"",D{i}-E{i})'; ws[f'G{i}'] = f'=IF(C{i}=0,"",D{i}/C{i})'; ws[f'H{i}'] = f'=IF(C{i}=0,"",E{i}/C{i})'
    ws[f'I{i}'] = f'=COUNTIFS({crit},{R_},"Pending")'
    ws[f'E{i}'].number_format = '0.00'; ws[f'F{i}'].number_format = '+0.00;-0.00;0.00'; ws[f'G{i}'].number_format = PCT; ws[f'H{i}'].number_format = PCT
hdr = ['Date', 'Market', 'Graded', 'Hits', 'Expected hits', 'Hits − expected', 'Hit rate', 'Avg P(hit)', 'Pending']
ws['A4'] = 'Record by market'; ws['A4'].font = bold
mkts = ['Goal', '1+ point', '3+ SOG']
table(ws, 5, hdr, [['All', m] + [None] * 7 for m in mkts], [12, 10, 9, 7, 13, 15, 10, 11, 9])
for k, m in enumerate(mkts): stats_row(ws, 6 + k, f'{M_},B{6+k}')
ws['A10'] = 'By night'; ws['A10'].font = bold
pairs = sorted({(t['date'], t['market']) for t in tracker}, key=lambda x: (x[0], mkts.index(x[1]) if x[1] in mkts else 9), reverse=True)
table(ws, 11, hdr, [[d, m] + [None] * 7 for d, m in pairs])
for k in range(len(pairs)):
    i = 12 + k; stats_row(ws, i, f'{A_},A{i},{M_},B{i}')
# COUNTIFS with two criteria pairs: rebuild with both conditions
for k in range(len(pairs)):
    i = 12 + k; c2 = f'{A_},A{i},{M_},B{i}'
    ws[f'C{i}'] = f'=COUNTIFS({c2},{R_},"Hit")+COUNTIFS({c2},{R_},"Miss")'
    ws[f'D{i}'] = f'=COUNTIFS({c2},{R_},"Hit")'
    ws[f'E{i}'] = f'=SUMIFS({P_},{c2},{R_},"Hit")+SUMIFS({P_},{c2},{R_},"Miss")'
    ws[f'I{i}'] = f'=COUNTIFS({c2},{R_},"Pending")'
ws[f'A{13+len(pairs)}'] = 'A few nights tell you almost nothing — judge each market after a few hundred graded picks.'; ws[f'A{13+len(pairs)}'].font = muted

# ---------------- Parlay tracker
ws = wb.create_sheet('Parlay tracker')
header_block(ws, 'Parlay tracker', 'Graded the next morning. A leg whose player did not play is voided and the parlay is graded on the rest, as most books do.')
PT0 = 4
last = table(ws, PT0, ['Date', 'Type', 'Legs', 'P(all hit)', 'Fair odds', 'Model', 'Result', 'Won'],
             [[p['date'], p['type'], p['parlay'], p['p_parlay'], None, p['model'], p['result'], None] for p in plog],
             [12, 18, 80, 11, 11, 8, 20, 6])
for i in range(PT0 + 1, last + 1):
    ws[f'D{i}'].number_format = PCT; ws[f'E{i}'] = FAIR.format(c=f'D{i}')
    ws[f'H{i}'] = f'=IF(LEFT(G{i},3)="Won",1,IF(LEFT(G{i},4)="Lost",0,""))'
if last > PT0:
    rng = f'G{PT0+1}:G{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(G{PT0+1},3)="Won"'], fill=green))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(G{PT0+1},4)="Lost"'], fill=red))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Pending"'], fill=grey))
ws.freeze_panes = f'C{PT0+1}'
PTL = max(last, PT0 + 1)

# ---------------- Parlay results
ws = wb.create_sheet('Parlay results')
header_block(ws, 'Parlay results', 'Expected wins = sum of the model\'s parlay probabilities. Low-probability parlays need many nights before the record means anything.')
types = [b['parlay'] for b in pbt] or ['Goals 2-leg', 'Goals 3-leg', 'Goals 4-leg', 'Points 2-leg', 'Points 3-leg', 'Shots 3+ 2-leg',
                                       'Shots 3+ 3-leg', 'Shots 2+ 3-leg', 'Shots 2+ 4-leg', 'Goals teammates']
table(ws, 4, ['Type', 'Graded', 'Won', 'Expected wins', 'Won − expected', 'Win rate', 'Pending'], [[t] + [None] * 6 for t in types],
      [18, 9, 7, 14, 15, 10, 9])
Ty, Rs, Pp = f"'Parlay tracker'!$B${PT0+1}:$B${PTL}", f"'Parlay tracker'!$G${PT0+1}:$G${PTL}", f"'Parlay tracker'!$D${PT0+1}:$D${PTL}"
for k in range(len(types)):
    i = 5 + k
    ws[f'B{i}'] = f'=COUNTIFS({Ty},A{i},{Rs},"Won*")+COUNTIFS({Ty},A{i},{Rs},"Lost*")'
    ws[f'C{i}'] = f'=COUNTIFS({Ty},A{i},{Rs},"Won*")'
    ws[f'D{i}'] = f'=SUMIFS({Pp},{Ty},A{i},{Rs},"Won*")+SUMIFS({Pp},{Ty},A{i},{Rs},"Lost*")'
    ws[f'E{i}'] = f'=IF(B{i}=0,"",C{i}-D{i})'; ws[f'F{i}'] = f'=IF(B{i}=0,"",C{i}/B{i})'
    ws[f'G{i}'] = f'=COUNTIFS({Ty},A{i},{Rs},"Pending")'
    ws[f'D{i}'].number_format = '0.00'; ws[f'E{i}'].number_format = '+0.00;-0.00;0.00'; ws[f'F{i}'].number_format = PCT
r = 6 + len(types)
ws[f'A{r}'] = 'Expected wins use the full parlay probability; when a leg is voided the true chance was higher, so treat this as approximate.'; ws[f'A{r}'].font = muted
ws[f'A{r+2}'] = 'Backtest — 2025-26 season (models trained on 2024-25 only), one parlay of each type per night'; ws[f'A{r+2}'].font = bold
e = table(ws, r + 3, ['Type', 'Nights', 'Model predicted', 'Actually hit', 'Gap'],
          [[b['parlay'], b['nights'], b['predicted_pct'] / 100, b['actual_pct'] / 100, None] for b in pbt])
for i in range(r + 4, e + 1):
    ws[f'C{i}'].number_format = PCT; ws[f'D{i}'].number_format = PCT
    ws[f'E{i}'] = f'=D{i}-C{i}'; ws[f'E{i}'].number_format = PCTD

# ---------------- Model changes
ws = wb.create_sheet('Model changes')
header_block(ws, 'Model improvement log', 'Each morning new ideas are tested per market. One is adopted only if it lowers log loss on BOTH test windows (A: 2025-26 season; B: most recent 60 game-days).')
table(ws, 4, ['Date', 'Market', 'Idea', 'Result', 'Model', 'Model log loss A', 'Idea log loss A', 'Model log loss B', 'Idea log loss B'],
      [[c['date'], c.get('market_label', c.get('market', '')), c['idea'], c['result'], c['champion'], c.get('champ_A'), c.get('chal_A'),
        c.get('champ_B'), c.get('chal_B')] for c in changes], [12, 16, 40, 20, 8, 14, 14, 14, 14])
for row in ws.iter_rows(min_row=5, min_col=6, max_col=9):
    for c in row: c.number_format = '0.00000'
ws.freeze_panes = 'D5'

# ---------------- Model accuracy
ws = wb.create_sheet('Model accuracy')
header_block(ws, 'How accurate are the models?', 'Lower log loss / Brier is better; higher AUC is better. All on games the models were not trained on.')
ws['A4'] = 'Current models vs a recent-form baseline — most recent 60 game-days'; ws['A4'].font = bold
e = table(ws, 5, ['Market', 'Model', 'Log loss', 'AUC', 'Brier', 'Avg predicted', 'Actually hit', 'Player-games'],
          [[a['market'], a['model'], a['log_loss'], a['auc'], a['brier'], a['avg_pred_pct'] / 100, a['actual_pct'] / 100, a['games']] for a in acc],
          [18, 24, 11, 9, 10, 13, 13, 13])
for i in range(6, e + 1):
    ws[f'C{i}'].number_format = '0.00000'; ws[f'D{i}'].number_format = '0.000'; ws[f'E{i}'].number_format = '0.0000'
    ws[f'F{i}'].number_format = PCT; ws[f'G{i}'].number_format = PCT; ws[f'H{i}'].number_format = '#,##0'
r = e + 2; ws[f'A{r}'] = 'How often the top picks hit — 2025-26 (models trained on 2024-25 only)'; ws[f'A{r}'].font = bold
e = table(ws, r + 1, ['Market', 'Picks', 'Predicted', 'Actual', 'Picks graded'],
          [[t['market'], t['picks'], t['predicted_pct'] / 100, t['actual_pct'] / 100, t['picks_n']] for t in top])
for i in range(r + 2, e + 1):
    ws[f'C{i}'].number_format = PCT; ws[f'D{i}'].number_format = PCT; ws[f'E{i}'].number_format = '#,##0'
r = e + 2; ws[f'A{r}'] = 'Goal model calibration — 2025-26 season'; ws[f'A{r}'].font = bold
e = table(ws, r + 1, ['Predicted band', 'Avg predicted', 'Actually scored', 'Gap', 'Player-games'],
          [[c['band'], c['predicted_pct'] / 100, c['actual_pct'] / 100, None, c['n']] for c in cal])
for i in range(r + 2, e + 1):
    ws[f'B{i}'].number_format = PCT; ws[f'C{i}'].number_format = PCT
    ws[f'D{i}'] = f'=C{i}-B{i}'; ws[f'D{i}'].number_format = PCTD; ws[f'E{i}'].number_format = '#,##0'
ws[f'A{e+2}'] = 'Source: MoneyPuck game-by-game and shot data (moneypuck.com); grading from NHL boxscores (api-web.nhle.com).'; ws[f'A{e+2}'].font = muted

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
    for row in s.iter_rows():
        for c in row:
            if c.font and c.font.name != F: c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.save(PATH)
print('saved', PATH)
